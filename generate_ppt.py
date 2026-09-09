from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = "毕业生就业信息管理系统_演示.pptx"
W, H = 13.333, 7.5
NAVY = RGBColor(18, 45, 79)
BLUE = RGBColor(25, 118, 210)
TEAL = RGBColor(0, 150, 136)
ORANGE = RGBColor(245, 124, 0)
BG = RGBColor(245, 248, 251)
TEXT = RGBColor(37, 50, 60)
MUTED = RGBColor(96, 125, 139)
WHITE = RGBColor(255, 255, 255)
LIGHT_BLUE = RGBColor(225, 239, 252)
LIGHT_TEAL = RGBColor(224, 242, 241)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = color


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=20, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         font="Microsoft YaHei", valign=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return tb


def title(slide, heading, subtitle=None):
    text(slide, heading, 0.65, 0.38, 8.8, 0.55, 27, NAVY, True)
    box(slide, 0.65, 1.04, 0.72, 0.06, BLUE)
    if subtitle:
        text(slide, subtitle, 1.55, 0.98, 10.8, 0.32, 10, MUTED)


def footer(slide, page):
    text(slide, "毕业生就业信息管理系统 · 课程设计演示", 0.65, 7.12, 7, 0.2, 9, MUTED)
    text(slide, f"{page:02d}", 12.1, 7.08, 0.55, 0.25, 10, BLUE, True, PP_ALIGN.RIGHT)


def bullet_list(slide, items, x, y, w, line_h=0.52, size=17, color=TEXT):
    for i, item in enumerate(items):
        yy = y + i * line_h
        box(slide, x, yy + 0.13, 0.12, 0.12, BLUE, BLUE, True)
        text(slide, item, x + 0.25, yy, w - 0.25, line_h, size, color)


def pill(slide, value, x, y, w, fill, color=WHITE):
    box(slide, x, y, w, 0.34, fill, fill, True)
    text(slide, value, x, y + 0.01, w, 0.28, 11, color, True, PP_ALIGN.CENTER)


def add_slide(prs, heading, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide); title(slide, heading, subtitle)
    footer(slide, len(prs.slides)); return slide


prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)

# 1 cover
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide, NAVY)
box(slide, 0, 0, W, H, NAVY)
box(slide, 0.75, 1.05, 0.12, 4.95, TEAL)
text(slide, "毕业生就业信息管理系统", 1.2, 1.45, 10.8, 0.75, 34, WHITE, True)
text(slide, "基于 Python Flask + SQLite 的课程设计演示", 1.23, 2.35, 9, 0.42, 19, RGBColor(190, 220, 242))
pill(slide, "学生端", 1.23, 3.45, 1.15, BLUE); pill(slide, "审核端", 2.55, 3.45, 1.15, TEAL); pill(slide, "统计端", 3.87, 3.45, 1.15, ORANGE)
text(slide, "信息录入 · 权限控制 · 协议审核 · 就业统计", 1.23, 4.25, 8.8, 0.35, 17, WHITE)
text(slide, "课程设计答辩材料", 1.23, 6.55, 5, 0.3, 12, RGBColor(190, 220, 242))
text(slide, "2026", 11.2, 6.55, 1.1, 0.3, 12, RGBColor(190, 220, 242), False, PP_ALIGN.RIGHT)

# 2 background
slide = add_slide(prs, "一、项目背景与建设目标", "面向毕业生就业协议收集、审核和统计的轻量化管理工具")
box(slide, 0.7, 1.45, 5.7, 4.75, WHITE, WHITE, True)
text(slide, "现实问题", 1.05, 1.78, 2.2, 0.35, 20, NAVY, True)
bullet_list(slide, ["就业信息分散在纸质表格或多个文件中", "学生提交、老师审核缺少统一流程", "就业数据难以实时汇总和按条件查询", "人工统计容易出现重复、遗漏和版本不一致"], 1.05, 2.35, 4.8, 0.65, 17)
box(slide, 6.75, 1.45, 5.85, 4.75, LIGHT_BLUE, LIGHT_BLUE, True)
text(slide, "建设目标", 7.1, 1.78, 2.2, 0.35, 20, NAVY, True)
bullet_list(slide, ["学生在线维护个人就业信息", "就业办集中审核并留下意见", "按状态、姓名、学号快速查询", "自动生成就业状态与单位性质统计"], 7.1, 2.35, 4.8, 0.65, 17)

# 3 requirements
slide = add_slide(prs, "二、系统需求与角色权限", "围绕“录入—审核—统计”形成闭环")
cards = [(0.8, "学生", BLUE, ["注册/登录", "查看个人信息", "录入或修改就业信息", "查看审核状态"]),
         (4.55, "就业办工作人员", TEAL, ["登录后台", "查询全部记录", "审核通过/驳回", "查看统计结果"]),
         (8.3, "系统管理员能力", ORANGE, ["SQLite 数据持久化", "角色权限隔离", "表单必填校验", "演示数据初始化"])]
for x, head, color, items in cards:
    box(slide, x, 1.55, 3.2, 4.9, WHITE, WHITE, True)
    box(slide, x, 1.55, 3.2, 0.75, color, color, True)
    text(slide, head, x + 0.2, 1.71, 2.8, 0.35, 19, WHITE, True, PP_ALIGN.CENTER)
    bullet_list(slide, items, x + 0.35, 2.65, 2.65, 0.72, 16)

# 4 architecture
slide = add_slide(prs, "三、系统总体架构", "采用简单、易理解、适合课程演示的 Flask 单体架构")
box(slide, 0.9, 1.55, 11.5, 0.85, WHITE, WHITE, True)
text(slide, "浏览器 / Web 页面", 1.05, 1.8, 2.5, 0.3, 18, NAVY, True, PP_ALIGN.CENTER)
text(slide, "→", 3.95, 1.75, 0.5, 0.35, 25, BLUE, True, PP_ALIGN.CENTER)
text(slide, "Flask 路由与权限控制", 4.55, 1.8, 3, 0.3, 18, NAVY, True, PP_ALIGN.CENTER)
text(slide, "→", 7.85, 1.75, 0.5, 0.35, 25, BLUE, True, PP_ALIGN.CENTER)
text(slide, "SQLAlchemy 数据模型", 8.45, 1.8, 2.8, 0.3, 18, NAVY, True, PP_ALIGN.CENTER)
box(slide, 2.1, 3.05, 2.55, 1.35, LIGHT_BLUE, LIGHT_BLUE, True); text(slide, "Jinja2 模板\n学生端 / 工作人员端", 2.25, 3.28, 2.25, 0.8, 17, BLUE, True, PP_ALIGN.CENTER)
box(slide, 5.35, 3.05, 2.55, 1.35, LIGHT_TEAL, LIGHT_TEAL, True); text(slide, "认证与权限\nSession + 角色装饰器", 5.5, 3.28, 2.25, 0.8, 17, TEAL, True, PP_ALIGN.CENTER)
box(slide, 8.6, 3.05, 2.55, 1.35, RGBColor(255, 243, 224), RGBColor(255, 243, 224), True); text(slide, "SQLite 数据库\nUser / Graduate / Employment", 8.75, 3.28, 2.25, 0.8, 16, ORANGE, True, PP_ALIGN.CENTER)
text(slide, "特点：依赖少、结构清晰、启动简单，适合课堂演示和功能讲解", 2.0, 5.35, 9.4, 0.4, 18, NAVY, True, PP_ALIGN.CENTER)

# 5 database
slide = add_slide(prs, "四、数据库设计", "三张核心表覆盖用户、毕业生和就业记录")
tables = [(0.8, "User 用户表", BLUE, ["id · username", "password_hash", "role", "graduate_id"]),
          (4.55, "Graduate 毕业生表", TEAL, ["id · student_no", "name · major", "class_name", "phone"]),
          (8.3, "Employment 就业表", ORANGE, ["employment_date", "employer · position", "archive_destination", "status · review_comment"])]
for x, head, color, rows in tables:
    box(slide, x, 1.65, 3.2, 3.8, WHITE, RGBColor(220, 228, 235), True)
    box(slide, x, 1.65, 3.2, 0.62, color, color, True)
    text(slide, head, x + 0.12, 1.8, 2.95, 0.3, 16, WHITE, True, PP_ALIGN.CENTER)
    for i, row in enumerate(rows):
        text(slide, row, x + 0.25, 2.48 + i * 0.58, 2.7, 0.35, 15, TEXT, i == 0)
text(slide, "1 个毕业生最多对应 1 条就业记录", 3.1, 6.05, 7.1, 0.35, 19, NAVY, True, PP_ALIGN.CENTER)

# 6 flow
slide = add_slide(prs, "五、核心业务流程", "学生提交后进入待审核，工作人员处理后形成最终状态")
steps = [(0.9, "注册 / 登录", BLUE, "账号与学籍信息"), (3.25, "填写就业信息", TEAL, "单位、时间、档案地"), (5.6, "提交审核", ORANGE, "状态：待审核"), (7.95, "工作人员审核", NAVY, "通过或驳回"), (10.3, "统计查询", TEAL, "按状态汇总")]
for i, (x, head, color, sub) in enumerate(steps):
    box(slide, x, 2.1, 1.85, 1.45, color, color, True)
    text(slide, str(i + 1), x + 0.08, 2.26, 0.35, 0.35, 18, WHITE, True, PP_ALIGN.CENTER)
    text(slide, head, x + 0.1, 2.68, 1.65, 0.3, 15, WHITE, True, PP_ALIGN.CENTER)
    text(slide, sub, x + 0.1, 3.05, 1.65, 0.28, 10, WHITE, False, PP_ALIGN.CENTER)
    if i < len(steps) - 1: text(slide, "→", x + 1.93, 2.57, 0.38, 0.35, 25, MUTED, True, PP_ALIGN.CENTER)
box(slide, 2.05, 4.55, 9.4, 0.9, WHITE, WHITE, True)
text(slide, "状态流转：", 2.35, 4.82, 1.35, 0.3, 17, NAVY, True)
pill(slide, "待审核", 3.85, 4.78, 1.15, ORANGE); text(slide, "→", 5.1, 4.77, 0.35, 0.3, 21, MUTED, True, PP_ALIGN.CENTER); pill(slide, "已通过", 5.55, 4.78, 1.15, TEAL); text(slide, "或", 6.85, 4.77, 0.35, 0.3, 14, MUTED, False, PP_ALIGN.CENTER); pill(slide, "已驳回", 7.3, 4.78, 1.15, RGBColor(211, 47, 47))

# 7 student
slide = add_slide(prs, "六、学生端功能演示", "登录后只看到与本人相关的数据")
box(slide, 0.85, 1.5, 5.65, 4.9, WHITE, WHITE, True)
text(slide, "学生个人中心", 1.15, 1.83, 2.7, 0.35, 21, NAVY, True)
pill(slide, "student001", 4.7, 1.84, 1.25, BLUE)
for i, (k, v) in enumerate([("学号", "2026001"), ("姓名", "张三"), ("专业", "计算机科学与技术"), ("班级", "计科2201班"), ("电话", "13800000000")]):
    y = 2.48 + i * 0.56; text(slide, k, 1.2, y, 0.75, 0.3, 15, MUTED); text(slide, v, 2.1, y, 3.75, 0.3, 15, TEXT, True)
box(slide, 6.8, 1.5, 5.65, 4.9, LIGHT_BLUE, LIGHT_BLUE, True)
text(slide, "就业信息维护", 7.1, 1.83, 2.7, 0.35, 21, NAVY, True)
bullet_list(slide, ["就业时间、就业单位", "单位性质、工作岗位", "档案发往地、备注", "修改后自动重新进入待审核"], 7.15, 2.5, 4.55, 0.68, 17)
pill(slide, "提交审核", 7.15, 5.35, 1.4, BLUE)

# 8 staff
slide = add_slide(prs, "七、工作人员端功能演示", "集中完成查询、审核和统计")
box(slide, 0.85, 1.55, 3.55, 4.85, WHITE, WHITE, True); text(slide, "记录查询", 1.15, 1.88, 2.5, 0.35, 20, NAVY, True); bullet_list(slide, ["姓名 / 学号关键词", "待审核 / 已通过 / 已驳回", "查看完整就业信息"], 1.15, 2.55, 2.8, 0.75, 16)
box(slide, 4.9, 1.55, 3.55, 4.85, LIGHT_TEAL, LIGHT_TEAL, True); text(slide, "审核处理", 5.2, 1.88, 2.5, 0.35, 20, TEAL, True); bullet_list(slide, ["核对就业协议相关信息", "填写审核意见", "选择通过或驳回", "记录审核时间"], 5.2, 2.55, 2.8, 0.75, 16)
box(slide, 8.95, 1.55, 3.55, 4.85, RGBColor(255, 243, 224), RGBColor(255, 243, 224), True); text(slide, "数据统计", 9.25, 1.88, 2.5, 0.35, 20, ORANGE, True); bullet_list(slide, ["记录总数", "待审核数量", "审核通过数量", "单位性质分布"], 9.25, 2.55, 2.8, 0.75, 16)

# 9 demo/run
slide = add_slide(prs, "八、运行方式与演示账号", "从安装到启动只需要几条命令")
box(slide, 0.9, 1.55, 7.05, 4.75, NAVY, NAVY, True)
text(slide, "启动命令", 1.25, 1.9, 2.3, 0.35, 20, WHITE, True)
commands = ["pip install -r requirements.txt", "flask --app app init-demo", "python app.py", "浏览器访问 http://127.0.0.1:5000"]
for i, cmd in enumerate(commands):
    text(slide, f"{i + 1:02d}", 1.3, 2.55 + i * 0.72, 0.45, 0.3, 14, RGBColor(130, 210, 200), True)
    text(slide, cmd, 1.95, 2.53 + i * 0.72, 5.45, 0.34, 16, WHITE, False, font="Consolas")
box(slide, 8.35, 1.55, 4.1, 4.75, WHITE, WHITE, True)
text(slide, "演示账号", 8.7, 1.9, 2.3, 0.35, 20, NAVY, True)
pill(slide, "学生", 8.7, 2.65, 0.9, BLUE); text(slide, "student001", 9.8, 2.65, 1.9, 0.3, 16, TEXT, True); text(slide, "123456", 9.8, 3.08, 1.9, 0.3, 15, MUTED)
pill(slide, "工作人员", 8.7, 4.0, 1.25, TEAL); text(slide, "admin", 10.1, 4.0, 1.5, 0.3, 16, TEXT, True); text(slide, "admin123", 10.1, 4.43, 1.5, 0.3, 15, MUTED)

# 10 summary
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide, NAVY)
text(slide, "项目总结", 0.9, 0.85, 4.5, 0.6, 31, WHITE, True)
text(slide, "用一个轻量系统，把就业信息管理流程串起来", 0.92, 1.65, 8.7, 0.42, 20, RGBColor(190, 220, 242))
for i, (head, sub, color) in enumerate([("可用", "功能闭环完整", BLUE), ("易懂", "结构清晰易讲解", TEAL), ("可扩展", "可接入真实业务", ORANGE)]):
    x = 1.0 + i * 4.05
    box(slide, x, 3.0, 3.25, 1.7, color, color, True)
    text(slide, head, x + 0.2, 3.28, 2.85, 0.42, 25, WHITE, True, PP_ALIGN.CENTER)
    text(slide, sub, x + 0.2, 3.92, 2.85, 0.3, 15, WHITE, False, PP_ALIGN.CENTER)
text(slide, "谢谢！", 0.9, 6.05, 4.0, 0.48, 23, WHITE, True)
text(slide, "Q&A", 11.1, 6.05, 1.35, 0.48, 23, RGBColor(190, 220, 242), True, PP_ALIGN.RIGHT)

prs.save(OUT)
print(OUT)
